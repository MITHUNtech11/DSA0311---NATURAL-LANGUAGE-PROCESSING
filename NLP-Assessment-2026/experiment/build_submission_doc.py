"""
Script to generate the final completed submission document:
DSA03-NLP_Lab_Experiments_1_to_5_Completed.docx

Strictly adheres to:
1. Original question pages preserved at the start.
2. Clean page break & prominent section divider.
3. Experiments 1 to 5 strictly formatted with Code and Output.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def read_text_file(filepath):
    """Safely reads file supporting utf-8 and utf-16 (from PowerShell Tee-Object)."""
    for enc in ["utf-8", "utf-16", "utf-16-le", "cp1252"]:
        try:
            with open(filepath, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(filepath, "r", errors="ignore") as f:
        return f.read()

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text, bg_color="F5F7FA"):
    """Adds a stylized code block within a single-cell table with light background."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border styling (thin border)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="0969DA"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    lines = code_text.strip().split('\n')
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9.0)
        run.font.color.rgb = RGBColor(36, 41, 47)
    
    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def add_output_block(doc, output_text, bg_color="F8F9FA"):
    """Adds a stylized terminal output block within a single-cell table."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    
    cell = tbl.cell(0, 0)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border styling (terminal dark green border)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:left w:val="single" w:sz="12" w:space="0" w:color="2DA44E"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="D0D7DE"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    
    lines = output_text.strip().split('\n')
    for idx, line in enumerate(lines):
        if idx > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(20, 24, 30)
        
    # Add spacing after table
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(6)

def build_completed_document():
    base_dir = r"C:\Users\mithu\OneDrive\Desktop\DSA0311 - NATURAL LANGUAGE PROCESSING\NLP-Assessment-2026\experiment"
    orig_docx = os.path.join(base_dir, "lab_program_list(1).docx")
    out_docx = os.path.join(base_dir, "DSA03-NLP_Lab_Experiments_1_to_5_Completed.docx")

    # Step 1: Open original document preserving original question page
    print(f"Loading original document: {orig_docx}")
    doc = docx.Document(orig_docx)

    # Step 2: Insert Clean Page Break and Prominent Section Divider
    doc.add_page_break()

    divider_p = doc.add_paragraph()
    divider_p.paragraph_format.space_before = Pt(18)
    divider_p.paragraph_format.space_after = Pt(6)
    divider_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div = divider_p.add_run("==================================================================")
    r_div.font.name = "Arial"
    r_div.font.size = Pt(11)
    r_div.font.bold = True
    r_div.font.color.rgb = RGBColor(31, 78, 121)

    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("DSA0311 - NATURAL LANGUAGE PROCESSING")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(31, 78, 121)

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_before = Pt(0)
    subtitle_p.paragraph_format.space_after = Pt(12)
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = subtitle_p.add_run("LABORATORY PRACTICAL RECORD: EXPERIMENTS 1 TO 5")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(80, 80, 80)

    divider_p2 = doc.add_paragraph()
    divider_p2.paragraph_format.space_before = Pt(0)
    divider_p2.paragraph_format.space_after = Pt(24)
    divider_p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_div2 = divider_p2.add_run("==================================================================")
    r_div2.font.name = "Arial"
    r_div2.font.size = Pt(11)
    r_div2.font.bold = True
    r_div2.font.color.rgb = RGBColor(31, 78, 121)

    # Experiments metadata
    experiments = [
        {
            "num": 1,
            "title": "Regular Expressions in Python for Text Pattern Matching and Searching",
            "prompt": "Write program demonstrates how to use regular expressions in Python to match and search for patterns in text.",
            "dir": "Experiment-01"
        },
        {
            "num": 2,
            "title": "Finite State Automaton (FSA) for Pattern Recognition (Ending with 'ab')",
            "prompt": "Implement a basic finite state automaton that recognizes a specific language or pattern. In this example, we'll create a simple automaton to match strings ending with 'ab' using python.",
            "dir": "Experiment-02"
        },
        {
            "num": 3,
            "title": "Morphological Analysis using the NLTK Library in Python",
            "prompt": "Write program demonstrates how to perform morphological analysis using the NLTK library in Python.",
            "dir": "Experiment-03"
        },
        {
            "num": 4,
            "title": "Finite-State Machine (FSM) for Morphological Parsing of Noun Plurals",
            "prompt": "Implement a finite-state machine for morphological parsing. In this example, we'll create a simple machine to generate plural forms of English nouns using python.",
            "dir": "Experiment-04"
        },
        {
            "num": 5,
            "title": "Word Stemming using Porter Stemmer Algorithm",
            "prompt": "Use the Porter Stemmer algorithm to perform word stemming on a list of words using python libraries.",
            "dir": "Experiment-05"
        }
    ]

    for exp in experiments:
        exp_dir = os.path.join(base_dir, exp["dir"])
        code_file = os.path.join(exp_dir, "main.py")
        output_file = os.path.join(exp_dir, "output.txt")
        
        code_content = read_text_file(code_file)
        output_content = read_text_file(output_file)

        # Add Experiment Heading
        exp_h = doc.add_paragraph()
        exp_h.paragraph_format.space_before = Pt(16)
        exp_h.paragraph_format.space_after = Pt(4)
        exp_h.paragraph_format.keep_with_next = True
        r_h = exp_h.add_run(f"Experiment {exp['num']}: {exp['title']}")
        r_h.font.name = "Arial"
        r_h.font.size = Pt(13)
        r_h.font.bold = True
        r_h.font.color.rgb = RGBColor(31, 78, 121)

        # Problem Statement
        stmt_p = doc.add_paragraph()
        stmt_p.paragraph_format.space_before = Pt(0)
        stmt_p.paragraph_format.space_after = Pt(10)
        stmt_p.paragraph_format.keep_with_next = True
        r_stmt_lbl = stmt_p.add_run("Problem Statement: ")
        r_stmt_lbl.font.name = "Arial"
        r_stmt_lbl.font.bold = True
        r_stmt_lbl.font.size = Pt(10)
        r_stmt_lbl.font.color.rgb = RGBColor(50, 50, 50)
        r_stmt_txt = stmt_p.add_run(exp["prompt"])
        r_stmt_txt.font.name = "Arial"
        r_stmt_txt.font.italic = True
        r_stmt_txt.font.size = Pt(10)
        r_stmt_txt.font.color.rgb = RGBColor(80, 80, 80)

        # Code Heading
        code_h = doc.add_paragraph()
        code_h.paragraph_format.space_before = Pt(6)
        code_h.paragraph_format.space_after = Pt(4)
        code_h.paragraph_format.keep_with_next = True
        r_ch = code_h.add_run("Code:")
        r_ch.font.name = "Arial"
        r_ch.font.size = Pt(11)
        r_ch.font.bold = True
        r_ch.font.color.rgb = RGBColor(31, 78, 121)

        # Code Block
        add_code_block(doc, code_content)

        # Output Heading
        out_h = doc.add_paragraph()
        out_h.paragraph_format.space_before = Pt(8)
        out_h.paragraph_format.space_after = Pt(4)
        out_h.paragraph_format.keep_with_next = True
        r_oh = out_h.add_run("Output:")
        r_oh.font.name = "Arial"
        r_oh.font.size = Pt(11)
        r_oh.font.bold = True
        r_oh.font.color.rgb = RGBColor(45, 122, 60)

        # Output Block
        add_output_block(doc, output_content)

        # Add page break between experiments (except after the last one)
        if exp["num"] < 5:
            doc.add_page_break()

    doc.save(out_docx)
    print(f"Successfully generated document: {out_docx}")

if __name__ == "__main__":
    build_completed_document()
